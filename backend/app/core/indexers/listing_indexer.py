"""All Listings + Category Listings SKU 索引器。"""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from app.core.parsers.sku import SKUInfo, parse_sku


def _normalize_header(value: Any) -> str:
    """标准化表头名称，兼容空格、大小写和 BOM。"""
    if value is None:
        return ""
    return str(value).replace("\ufeff", "").strip().casefold()


def _normalize_cell(value: Any) -> Any:
    """读取单元格时去掉字符串首尾空白。"""
    if isinstance(value, str):
        return value.strip()
    return value


def _is_missing_value(value: Any) -> bool:
    return value is None or value == ""


class ListingIndexer:
    """从 All Listings 与 Category Listings 构建 SKU 索引。"""

    def build_index(self, listings_files: list[Path], category_files: list[Path]) -> dict[str, Any]:
        all_rows: list[dict[str, Any]] = []
        for listings_file in listings_files:
            all_rows.extend(self._read_all_listings_report(listings_file))
        category_rows = self._read_category_listing_reports(category_files)

        by_sku: dict[str, dict[str, Any]] = {}
        by_prefix: dict[str, dict[str, dict[str, Any]]] = defaultdict(dict)
        unparsed_rows: list[dict[str, Any]] = []

        for row in all_rows:
            self._merge_row(
                by_sku=by_sku,
                by_prefix=by_prefix,
                row=row,
                source_key="all_listings_data",
            )
            if row["parsed_sku"] is None:
                unparsed_rows.append(row)

        for row in category_rows:
            self._merge_row(
                by_sku=by_sku,
                by_prefix=by_prefix,
                row=row,
                source_key="category_data",
            )
            if row["parsed_sku"] is None:
                unparsed_rows.append(row)

        return {
            "by_sku": by_sku,
            "by_prefix": dict(by_prefix),
            "all_listings_rows": all_rows,
            "category_rows": category_rows,
            "unparsed_rows": unparsed_rows,
            "stats": {
                "all_listings_count": len(all_rows),
                "category_count": len(category_rows),
                "sku_count": len(by_sku),
                "prefix_count": len(by_prefix),
                "unparsed_count": len(unparsed_rows),
            },
        }

    def _read_all_listings_report(self, listings_file: Path) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []

        with listings_file.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            if reader.fieldnames:
                reader.fieldnames = [_normalize_header(name) for name in reader.fieldnames]

            for line_number, raw_row in enumerate(reader, start=2):
                row = {
                    key: _normalize_cell(value)
                    for key, value in raw_row.items()
                    if key is not None
                }
                sku = str(row.get("seller-sku") or "").strip().upper()
                if not sku:
                    continue

                parsed_sku = self._parse_sku_safe(sku)
                rows.append(
                    {
                        "sku": sku,
                        "prefix": parsed_sku.product_code if parsed_sku else None,
                        "parsed_sku": parsed_sku,
                        "source": "all_listings",
                        "line_number": line_number,
                        "data": row,
                    }
                )

        return rows

    def _read_category_listing_reports(self, category_files: list[Path]) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []

        for file_path in category_files:
            workbook = load_workbook(file_path, data_only=True, read_only=True)
            try:
                worksheet, header_row_index, headers = self._find_category_sheet(workbook)

                for row_index, values in enumerate(
                    worksheet.iter_rows(min_row=header_row_index + 1, values_only=True),
                    start=header_row_index + 1,
                ):
                    if not any(value not in (None, "") for value in values):
                        continue

                    row = self._build_category_row(headers, values)
                    sku = str(row.get("sku") or row.get("seller sku") or row.get("seller-sku") or "").strip().upper()
                    if not sku:
                        continue

                    parsed_sku = self._parse_sku_safe(sku)
                    rows.append(
                        {
                            "sku": sku,
                            "prefix": parsed_sku.product_code if parsed_sku else None,
                            "parsed_sku": parsed_sku,
                            "source": "category",
                            "source_file": file_path.name,
                            "line_number": row_index,
                            "data": row,
                        }
                    )
            finally:
                workbook.close()

        return rows

    _TEMPLATE_SHEET_NAMES = frozenset({"template", "modèle", "vorlage", "modello", "plantilla"})

    def _find_category_sheet(self, workbook) -> tuple[Any, int, list[str]]:
        """优先选择 Template（含各语言翻译），否则选择第一个包含 SKU 列的 sheet。"""
        ordered_names = list(workbook.sheetnames)
        template_name = next(
            (name for name in ordered_names if name.casefold() in self._TEMPLATE_SHEET_NAMES),
            None,
        )
        if template_name is not None:
            ordered_names.remove(template_name)
            ordered_names.insert(0, template_name)

        for sheet_name in ordered_names:
            worksheet = workbook[sheet_name]
            scanned_rows = [
                (row_index, [_normalize_header(value) for value in values])
                for row_index, values in enumerate(
                    worksheet.iter_rows(min_row=1, max_row=min(25, worksheet.max_row), values_only=True),
                    start=1,
                )
            ]

            for row_index, headers in scanned_rows:
                if "contribution_sku#1.value" in headers:
                    return worksheet, row_index, headers

            for row_index, headers in scanned_rows:
                if any(header in ("sku", "seller sku", "seller-sku") for header in headers):
                    return worksheet, row_index, headers

        raise ValueError("No worksheet containing a 'SKU' column was found.")

    def _build_category_row(self, headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
        row: dict[str, Any] = {}
        bullet_point_count = 0
        generic_keyword_seen = False

        for header, value in zip(headers, values):
            if not header:
                continue

            normalized_value = _normalize_cell(value)

            bullet_point_key = self._bullet_point_key(header, next_index=bullet_point_count + 1)
            if bullet_point_key is not None:
                bullet_point_count = max(bullet_point_count, int(bullet_point_key.removeprefix("bullet_point")))
                row[bullet_point_key] = normalized_value
                continue

            if self._is_generic_keyword_header(header):
                if generic_keyword_seen:
                    continue
                row["generic_keywords"] = normalized_value
                generic_keyword_seen = True
                continue

            row[header] = normalized_value
            base_key = self._simplified_machine_header_key(header)
            if base_key and base_key not in row:
                row[base_key] = normalized_value

        return row

    def _bullet_point_key(self, header: str, *, next_index: int) -> str | None:
        if not (header.startswith("bullet_point") or header.startswith("bullet point")):
            return None

        match = re.search(r"#(\d+)", header)
        bullet_index = int(match.group(1)) if match else next_index
        return f"bullet_point{bullet_index}"

    def _is_generic_keyword_header(self, header: str) -> bool:
        return header in ("generic keywords", "generic keyword") or header.startswith("generic_keyword")

    def _simplified_machine_header_key(self, header: str) -> str | None:
        if "[" not in header and "#" not in header:
            return None
        if header.startswith("contribution_sku"):
            return "sku"

        base_key = re.split(r"[\[#]", header, maxsplit=1)[0].strip()
        return base_key or None

    def _merge_row(
        self,
        *,
        by_sku: dict[str, dict[str, Any]],
        by_prefix: dict[str, dict[str, dict[str, Any]]],
        row: dict[str, Any],
        source_key: str,
    ) -> None:
        sku = row["sku"]
        existing = by_sku.get(sku)

        if existing is None:
            existing = {
                "sku": sku,
                "prefix": row["prefix"],
                "parsed_sku": row["parsed_sku"],
                "all_listings_data": None,
                "category_data": None,
                "merged_data": {},
            }
            by_sku[sku] = existing

        if existing["parsed_sku"] is None and row["parsed_sku"] is not None:
            existing["parsed_sku"] = row["parsed_sku"]
            existing["prefix"] = row["prefix"]

        if source_key == "category_data" and existing["category_data"]:
            merged_category_data = dict(existing["category_data"])
            for key, value in row["data"].items():
                if key not in merged_category_data or _is_missing_value(merged_category_data[key]):
                    merged_category_data[key] = value
            existing["category_data"] = merged_category_data
        else:
            existing[source_key] = row["data"]

        merged_data = {}
        if existing["all_listings_data"]:
            merged_data.update(existing["all_listings_data"])
        if existing["category_data"]:
            merged_data.update(existing["category_data"])
        existing["merged_data"] = merged_data

        prefix = existing["prefix"]
        if prefix:
            by_prefix[prefix][sku] = existing

    def _parse_sku_safe(self, raw_sku: str) -> SKUInfo | None:
        try:
            return parse_sku(raw_sku)
        except ValueError:
            return None
