from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import sys

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.config import Country
import app.core.processors.add_color_size_processor as processor_module
from app.core.color_mapper import color_mapper
from app.core.indexers.listing_indexer import ListingIndexer
from app.core.processors.add_color_size_processor import AddColorSizeProcessor
from app.models.excel import ProcessRequest


def _sheet_rows(workbook_path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        headers = [str(value) if value is not None else "" for value in next(worksheet.iter_rows(max_row=1, values_only=True))]
        rows: list[dict[str, object]] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(dict(zip(headers, values)))
        return headers, rows
    finally:
        workbook.close()


def _template_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = next(name for name in workbook.sheetnames if name.lower() == "template")
        worksheet = workbook[sheet_name]
        headers = [
            str(worksheet.cell(row=2, column=column_index).value or "")
            for column_index in range(1, worksheet.max_column + 1)
        ]
        rows: list[dict[str, object]] = []
        for row_index in range(4, worksheet.max_row + 1):
            values = [worksheet.cell(row=row_index, column=column_index).value for column_index in range(1, worksheet.max_column + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(dict(zip(headers, values)))
        return headers, rows
    finally:
        workbook.close()


def _template_machine_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, object]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = next(name for name in workbook.sheetnames if name.lower() == "template")
        worksheet = workbook[sheet_name]
        headers = [
            str(worksheet.cell(row=3, column=column_index).value or "")
            for column_index in range(1, worksheet.max_column + 1)
        ]
        rows: list[dict[str, object]] = []
        for row_index in range(4, worksheet.max_row + 1):
            values = [worksheet.cell(row=row_index, column=column_index).value for column_index in range(1, worksheet.max_column + 1)]
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(dict(zip(headers, values)))
        return headers, rows
    finally:
        workbook.close()


class AddColorSizeProcessorTests(unittest.TestCase):
    def setUp(self) -> None:
        self._results_dir = tempfile.TemporaryDirectory()
        self._uploads_dir = tempfile.TemporaryDirectory()
        self._templates_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._results_dir.cleanup)
        self.addCleanup(self._uploads_dir.cleanup)
        self.addCleanup(self._templates_dir.cleanup)

        self._original_results_dir = processor_module.RESULTS_DIR
        self._original_uploads_dir = processor_module.UPLOADS_DIR
        self._original_templates_dir = processor_module.TEMPLATES_DIR
        self._original_mappings = color_mapper.mappings.copy()

        processor_module.RESULTS_DIR = Path(self._results_dir.name)
        processor_module.UPLOADS_DIR = Path(self._uploads_dir.name)
        processor_module.TEMPLATES_DIR = Path(self._templates_dir.name)
        color_mapper.mappings = {
            "BK": {"en": "Black", "fr": "Noir", "de": "Schwarz", "it": "Nero", "es": "Negro"},
            "BD": {"en": "Royal Blue", "fr": "Bleu Royal", "de": "Königsblau", "it": "Blu Reale", "es": "Azul Real"},
            "SG": {"en": "Sage Green", "fr": "", "de": "", "it": "", "es": ""},
        }

        self.addCleanup(self._restore_environment)

    def _restore_environment(self) -> None:
        processor_module.RESULTS_DIR = self._original_results_dir
        processor_module.UPLOADS_DIR = self._original_uploads_dir
        processor_module.TEMPLATES_DIR = self._original_templates_dir
        color_mapper.mappings = self._original_mappings

    def _real_uploads_dir(self) -> Path:
        return Path(__file__).resolve().parents[1] / "uploads"

    def _require_upload_files(self, *filenames: str) -> Path:
        uploads_dir = self._real_uploads_dir()
        missing = [name for name in filenames if not (uploads_dir / name).exists()]
        if missing:
            pytest.skip(f"Missing real upload files: {', '.join(missing)}")
        processor_module.UPLOADS_DIR = uploads_dir
        return uploads_dir

    def _build_real_index(self, *, listings_file: str, category_files: list[str]) -> dict[str, object]:
        uploads_dir = self._require_upload_files(listings_file, *category_files)
        return ListingIndexer().build_index(
            [uploads_dir / listings_file],
            [uploads_dir / filename for filename in category_files],
        )

    def _child_rows(self, rows: list[dict[str, object]]) -> list[dict[str, object]]:
        return [r for r in rows if r.get("Parentage") != "Parent" and r.get("parent_child") != "Parent"]

    def _create_template_workbook(self, country: Country = Country.UK) -> Path:
        template_path = processor_module.TEMPLATES_DIR / processor_module.COUNTRY_PROFILES[country].template_file
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Template"
        worksheet.append(["Metadata"])
        worksheet.append(
            [
                "Seller SKU",
                "Parent SKU",
                "Parentage",
                "Variation Theme",
                "Colour",
                "Colour Map",
                "Size",
                "Size Map",
                "Product Name",
                "Currency",
                "List Price with Tax for Display",
                "Main Image Url",
                "Other Image Url1",
                "Other Image Url2",
                "Other Image Url3",
                "Other Image Url4",
                "Other Image Url5",
                "external_product_id",
                "external_product_id_type",
                "bullet_point[marketplace_id=A1F83G8C2ARO7P][language_tag=en_GB]#1.value",
                "generic_keyword[marketplace_id=A1F83G8C2ARO7P][language_tag=en_GB]#1.value",
            ]
        )
        worksheet.append(
            [
                "item_sku",
                "parent_sku",
                "parent_child",
                "variation_theme",
                "color_name",
                "color_map",
                "size_name",
                "size_map",
                "item_name",
                "currency",
                "list_price_with_tax",
                "main_image_url",
                "other_image_url1",
                "other_image_url2",
                "other_image_url3",
                "other_image_url4",
                "other_image_url5",
                "external_product_id",
                "external_product_id_type",
                "bullet_point1",
                "generic_keywords",
            ]
        )
        worksheet.append(["stale", "stale", "stale"])
        workbook.save(template_path)
        workbook.close()
        return template_path

    def test_process_generates_grouped_workbook_for_listing_index_style_input(self) -> None:
        self._create_template_workbook()
        index = self._build_real_index(
            listings_file="UK_all_listings_20260313082242_All+Listings+Report_03-11-2026.txt",
            category_files=[
                "UK_category_20260313082255_0_APPAREL-DRESS.xlsm",
                "UK_category_20260313082258_1_DRESS.xlsm",
                "UK_category_20260313082300_2_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
        )

        result = AddColorSizeProcessor().process(
            index,
            country="UK",
            selected_prefixes=["EG02230"],
            target_colors=["BK"],
            start_size="04",
            end_size="04",
            size_step=1,
            output_filename="uk-output.xlsx",
        )

        self.assertEqual(result.output_file, "uk-output.xlsx")
        self.assertEqual(result.processed_count, 2)
        self.assertEqual(result.skipped_count, 0)
        self.assertEqual(result.errors, [])

        output_path = Path(self._results_dir.name) / result.output_file
        self.assertTrue(output_path.exists())

        workbook = load_workbook(output_path, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["Template"])
        finally:
            workbook.close()

        headers, template_rows = _template_rows(output_path)
        self.assertEqual(len(template_rows), 4)
        parent_rows = [row for row in template_rows if row.get("Parentage") == "Parent" or row.get("parent_child") == "Parent"]
        children = self._child_rows(template_rows)

        self.assertEqual(len(parent_rows), 2)
        self.assertEqual(len(children), 2)
        self.assertEqual([row["Seller SKU"] for row in parent_rows], ["EG02230", "EG02230"])
        self.assertEqual(
            sorted(str(row["Seller SKU"]) for row in children),
            ["EG02230BK04", "EG02230BK04-UK1"],
        )
        self.assertEqual(
            {str(row["Seller SKU"]): str(row["Parent SKU"]) for row in children},
            {
                "EG02230BK04": "EG02230",
                "EG02230BK04-UK1": "EG02230-UK1",
            },
        )
        self.assertTrue(all(row["Colour"] == "Black" for row in children))
        self.assertTrue(all(row["Colour Map"] == "Black" for row in children))
        self.assertTrue(all(row["Size"] == 8 for row in children))
        self.assertEqual(
            {str(row["Seller SKU"]): str(row["Other Image Url1"]) for row in children},
            {
                "EG02230BK04": "https://eppic.s3.amazonaws.com/EG02230BK-L2.jpg",
                "EG02230BK04-UK1": "https://eppic.s3.amazonaws.com/EG02230BK-L2.jpg",
            },
        )
        self.assertEqual(
            {str(row["Seller SKU"]): str(row["Other Image Url5"]) for row in children},
            {
                "EG02230BK04": "https://eppic.s3.amazonaws.com/EG02230BK-S-UK.jpg",
                "EG02230BK04-UK1": "https://eppic.s3.amazonaws.com/EG02230BK-S-UK.jpg",
            },
        )

    def _create_de_template_workbook(self) -> Path:
        template_path = processor_module.TEMPLATES_DIR / processor_module.COUNTRY_PROFILES[Country.DE].template_file
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "template"
        worksheet.append(["Metadata"])
        worksheet.append(
            [
                "Produkttyp",
                "Verkäufer-SKU",
                "Update / Löschen",
                "SKU des übergeordneten Produkts",
                "Variantenbestandteil",
                "Typ der untergeordneten Beziehung",
                "Varianten-Design",
                "Produktname",
                "Marke",
                "Hersteller-Barcode  ",
                "Barcode-Typ",
            ]
        )
        worksheet.append(
            [
                "feed_product_type",
                "item_sku",
                "update_delete",
                "parent_sku",
                "parent_child",
                "relationship_type",
                "variation_theme",
                "item_name",
                "brand_name",
                "external_product_id",
                "external_product_id_type",
            ]
        )
        worksheet.append(["stale"] * 11)
        workbook.save(template_path)
        workbook.close()
        return template_path

    def _create_it_template_workbook(self) -> Path:
        template_path = processor_module.TEMPLATES_DIR / processor_module.COUNTRY_PROFILES[Country.IT].template_file
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "template"
        worksheet.append(["Metadata"])
        worksheet.append(
            [
                "Seller SKU",
                "Parentage",
                "Update Delete",
                "Condition Type",
                "Colour",
                "Colour Map",
                "Country/Region Of Origin",
                "pattern_type",
                "supplier_declared_material_regulation1",
                "package_level",
                "package_weight",
                "package_weight_unit_of_measure",
                "package_length_unit_of_measure",
                "package_height_unit_of_measure",
                "package_width_unit_of_measure",
                "supplier_declared_dg_hz_regulation1",
            ]
        )
        worksheet.append(
            [
                "item_sku",
                "parent_child",
                "update_delete",
                "condition_type",
                "color_name",
                "color_map",
                "country_of_origin",
                "pattern_type",
                "supplier_declared_material_regulation1",
                "package_level",
                "package_weight",
                "package_weight_unit_of_measure",
                "package_length_unit_of_measure",
                "package_height_unit_of_measure",
                "package_width_unit_of_measure",
                "supplier_declared_dg_hz_regulation1",
            ]
        )
        worksheet.append(["stale"] * 16)
        workbook.save(template_path)
        workbook.close()
        return template_path

    def test_process_keeps_suffixless_eu_skus_in_output(self) -> None:
        self._create_de_template_workbook()
        index = self._build_real_index(
            listings_file="DE_all_listings_20260318062018_所有商品报告_03-18-2026.txt",
            category_files=[
                "DE_category_20260318062031_0_APPAREL_HEAD_NECK_COVERING-DRESS.xlsm",
                "DE_category_20260318062034_1_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
        )

        result = AddColorSizeProcessor().process(
            index,
            country="DE",
            selected_prefixes=["EA02333"],
            target_colors=["BK"],
            start_size="00",
            end_size="00",
            size_step=1,
            output_filename="de-suffixless-output.xlsx",
        )

        self.assertEqual(result.output_file, "de-suffixless-output.xlsx")
        self.assertEqual(result.processed_count, 1)
        self.assertEqual(result.skipped_count, 0)
        self.assertEqual(result.errors, [])

        output_path = Path(self._results_dir.name) / result.output_file
        _, rows = _template_machine_rows(output_path)
        children = self._child_rows(rows)

        self.assertEqual(len(rows), 2)
        self.assertEqual(len(children), 1)
        self.assertEqual(rows[0]["item_sku"], "EA02333")
        self.assertEqual(rows[0]["parent_child"], "Parent")
        self.assertEqual(children[0]["item_sku"], "EA02333BK00")
        self.assertEqual(children[0]["parent_sku"], "EA02333")

    def test_process_request_uses_listing_indexer_dict_output(self) -> None:
        self._create_template_workbook()
        self._require_upload_files(
            "UK_all_listings_20260313082242_All+Listings+Report_03-11-2026.txt",
            "UK_category_20260313082255_0_APPAREL-DRESS.xlsm",
            "UK_category_20260313082258_1_DRESS.xlsm",
            "UK_category_20260313082300_2_DRESS-UNDERGARMENT_SLIP.xlsm",
        )

        request = ProcessRequest(
            country=Country.UK,
            all_listings_files=["UK_all_listings_20260313082242_All+Listings+Report_03-11-2026.txt"],
            category_files=[
                "UK_category_20260313082255_0_APPAREL-DRESS.xlsm",
                "UK_category_20260313082258_1_DRESS.xlsm",
                "UK_category_20260313082300_2_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
            selected_prefixes=["EG02230"],
            target_colors=["BK"],
            start_size="04",
            end_size="04",
            size_step=1,
            mode="add-color",
        )

        result = AddColorSizeProcessor().process(request, output_filename="request-output")

        self.assertEqual(result.output_file, "request-output.xlsx")
        self.assertEqual(result.processed_count, 2)
        self.assertEqual(result.skipped_count, 0)

        output_path = Path(self._results_dir.name) / result.output_file
        _, rows = _template_rows(output_path)
        children = self._child_rows(rows)

        self.assertEqual(len(rows), 4)
        self.assertEqual(
            sorted(str(row["Seller SKU"]) for row in children),
            ["EG02230BK04", "EG02230BK04-UK1"],
        )
        self.assertEqual(
            {str(row["Seller SKU"]): str(row["Parent SKU"]) for row in children},
            {
                "EG02230BK04": "EG02230",
                "EG02230BK04-UK1": "EG02230-UK1",
            },
        )
        self.assertTrue(all(row["Colour"] == "Black" for row in children))

    def test_process_request_direct_sku_groups_variants_by_prefix(self) -> None:
        self._create_template_workbook()
        self._require_upload_files(
            "UK_all_listings_20260313082242_All+Listings+Report_03-11-2026.txt",
            "UK_category_20260313082255_0_APPAREL-DRESS.xlsm",
            "UK_category_20260313082258_1_DRESS.xlsm",
            "UK_category_20260313082300_2_DRESS-UNDERGARMENT_SLIP.xlsm",
        )

        request = ProcessRequest(
            country=Country.UK,
            all_listings_files=["UK_all_listings_20260313082242_All+Listings+Report_03-11-2026.txt"],
            category_files=[
                "UK_category_20260313082255_0_APPAREL-DRESS.xlsm",
                "UK_category_20260313082258_1_DRESS.xlsm",
                "UK_category_20260313082300_2_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
            input_mode="direct-sku",
            direct_skus=["EG02230BK04", "EG02230BD06", "EG02230BK04-UK1"],
            mode="add-color",
        )

        result = AddColorSizeProcessor().process(request, output_filename="direct-sku-output")

        self.assertEqual(result.output_file, "direct-sku-output.xlsx")
        self.assertEqual(result.processed_count, 4)
        self.assertEqual(result.skipped_count, 0)

        output_path = Path(self._results_dir.name) / result.output_file
        _, rows = _template_rows(output_path)
        children = self._child_rows(rows)

        self.assertEqual(len(rows), 6)
        self.assertEqual(
            sorted(str(row["Seller SKU"]) for row in children),
            ["EG02230BD06", "EG02230BD06-UK1", "EG02230BK04", "EG02230BK04-UK1"],
        )
        self.assertEqual(
            {str(row["Seller SKU"]): str(row["Parent SKU"]) for row in children},
            {
                "EG02230BK04": "EG02230",
                "EG02230BD06": "EG02230",
                "EG02230BK04-UK1": "EG02230-UK1",
                "EG02230BD06-UK1": "EG02230-UK1",
            },
        )

    def test_process_request_uses_fr_asin_for_de_external_product_id(self) -> None:
        self._create_de_template_workbook()
        uploads_dir = self._require_upload_files(
            "DE_all_listings_20260318062018_所有商品报告_03-18-2026.txt",
            "DE_category_20260318062031_0_APPAREL_HEAD_NECK_COVERING-DRESS.xlsm",
            "DE_category_20260318062034_1_DRESS-UNDERGARMENT_SLIP.xlsm",
        )
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            suffix=".txt",
            prefix="FR_all_listings_test_",
            dir=uploads_dir,
            delete=False,
        ) as handle:
            handle.write("seller-sku\tasin1\n")
            handle.write("EA02333BK00\tB0FRPANEU\n")
            fr_listings_name = Path(handle.name).name
        self.addCleanup(lambda: (uploads_dir / fr_listings_name).unlink(missing_ok=True))

        request = ProcessRequest(
            country=Country.DE,
            all_listings_files=["DE_all_listings_20260318062018_所有商品报告_03-18-2026.txt"],
            fr_all_listings_file=fr_listings_name,
            category_files=[
                "DE_category_20260318062031_0_APPAREL_HEAD_NECK_COVERING-DRESS.xlsm",
                "DE_category_20260318062034_1_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
            selected_prefixes=["EA02333"],
            target_colors=["BK"],
            start_size="00",
            end_size="00",
            size_step=1,
            mode="add-color",
        )

        result = AddColorSizeProcessor().process(request, output_filename="de-request-output")

        self.assertEqual(result.output_file, "de-request-output.xlsx")
        self.assertEqual(result.processed_count, 1)
        self.assertEqual(result.skipped_count, 0)

        output_path = Path(self._results_dir.name) / result.output_file
        _, rows = _template_machine_rows(output_path)
        children = self._child_rows(rows)

        self.assertEqual(len(rows), 2)
        self.assertEqual(len(children), 1)
        self.assertEqual(children[0]["item_sku"], "EA02333BK00")
        self.assertEqual(children[0]["parent_sku"], "EA02333")
        # Child rows must NOT carry external_product_id — only the parent row does
        self.assertIsNone(children[0].get("external_product_id"))
        self.assertIsNone(children[0].get("external_product_id_type"))
        parents = [r for r in rows if r.get("parent_child") == "Parent"]
        self.assertEqual(len(parents), 1)
        self.assertIsNotNone(parents[0].get("external_product_id"))

    def test_build_output_row_uses_de_package_fields_and_fixed_machine_values(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EA02333BK00",
                "parsed_sku": processor_module.parse_sku("EA02333BK00"),
                "price": 19.9,
                "asin": "B0TESTASIN",
                "all_listings_data": {"item_name": "Unterkleid Schwarz 34DE"},
                "category_data": {
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.length.value": 30.0,
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.width.value": 20.0,
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.height.value": 1.0,
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.length.unit": "Zentimeter",
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.width.unit": "Zentimeter",
                    "item_package_dimensions[marketplace_id=a1pa6795ukmfr9]#1.height.unit": "Zentimeter",
                    "item_package_weight[marketplace_id=a1pa6795ukmfr9]#1.value": 119.0,
                    "item_package_weight[marketplace_id=a1pa6795ukmfr9]#1.unit": "Gramm",
                },
            }
        )

        assert record is not None

        row = processor._build_output_row(
            source_record=record,
            new_color_code="BK",
            new_size_code="00",
            profile=processor_module.COUNTRY_PROFILES[Country.DE],
        )

        self.assertEqual(row["pattern_type"], "Uni")
        self.assertEqual(row["fulfillment_center_id"], "DEFAULT")
        self.assertEqual(row["is_adult_product"], "No")
        self.assertEqual(row["supplier_declared_dg_hz_regulation1"], "Nicht zutreffend")
        self.assertEqual(row["Colour"], "Schwarz")
        self.assertEqual(row["Colour Map"], "Schwarz")
        self.assertEqual(row["Package Length"], 30.0)
        self.assertEqual(row["Package Width"], 20.0)
        self.assertEqual(row["Package Height"], 1.0)
        self.assertEqual(row["Package Weight"], 119.0)
        self.assertEqual(row["package_length_unit_of_measure"], "centimeters")
        self.assertEqual(row["package_width_unit_of_measure"], "centimeters")
        self.assertEqual(row["package_height_unit_of_measure"], "centimeters")
        self.assertEqual(row["package_weight_unit_of_measure"], "kilograms")
        self.assertNotIn("Package Length Unit Of Measure", row)
        self.assertNotIn("Package Dimensions Unit Of Measure", row)
        self.assertNotIn("Package Weight Unit Of Measure", row)

    def test_build_output_row_keeps_new_sku_for_model_fields(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EE50126WH04-UK1",
                "parsed_sku": processor_module.parse_sku("EE50126WH04-UK1"),
                "all_listings_data": {
                    "item_name": "Slip White 8UK",
                    "model": "EE50126WH04-UK1",
                },
                "category_data": {
                    "model": "EE50126WH04-UK1",
                    "model[marketplace_id=a1f83g8c2aro7p]#1.value": "EE50126WH04-UK1",
                    "part_number": "EE50126WH04-UK1",
                },
            }
        )

        assert record is not None

        row = processor._build_output_row(
            source_record=record,
            new_color_code="BK",
            new_size_code="04",
            profile=processor_module.COUNTRY_PROFILES[Country.UK],
        )

        self.assertEqual(row["Seller SKU"], "EE50126BK04-UK1")
        self.assertEqual(row["Parent SKU"], "EE50126-UK1")
        self.assertEqual(row["Model Name"], "EE50126BK04-UK1")
        self.assertEqual(row["Model Number"], "EE50126BK04-UK1")
        self.assertEqual(row["model"], "EE50126BK04-UK1")
        self.assertEqual(row["part_number"], "EE50126BK04-UK1")

    def test_build_parent_row_copies_item_name_to_machine_field(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EE50126WH04-UK1",
                "parsed_sku": processor_module.parse_sku("EE50126WH04-UK1"),
                "all_listings_data": {"item_name": "Slip White 8UK"},
                "category_data": {"product type": "dress", "brand name": "Test Brand"},
            }
        )

        assert record is not None

        row = processor._build_parent_row(
            source_record=record,
            profile=processor_module.COUNTRY_PROFILES[Country.UK],
        )

        self.assertEqual(row["Seller SKU"], "EE50126")
        self.assertEqual(row["Product Name"], "Slip White 8UK")
        self.assertEqual(row["item_name"], "Slip White 8UK")

    def test_build_output_row_uses_it_static_fields_and_marketplace_aliases(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EA02333BK00",
                "parsed_sku": processor_module.parse_sku("EA02333BK00"),
                "price": 19.9,
                "asin": "B0TESTASIN",
                "all_listings_data": {"item_name": "Sottoveste Nero 32IT"},
                "category_data": {
                    "brand[marketplace_id=apj6jra9ng5v4][language_tag=it_it]#1.value": "Marchio Test",
                    "material[marketplace_id=apj6jra9ng5v4][language_tag=it_it]#1.value": "Poliestere",
                    "generic_keyword[marketplace_id=apj6jra9ng5v4][language_tag=it_it]#1.value": "sottoveste donna",
                    "item_package_dimensions[marketplace_id=apj6jra9ng5v4]#1.length.value": 30.0,
                    "item_package_dimensions[marketplace_id=apj6jra9ng5v4]#1.width.value": 20.0,
                    "item_package_dimensions[marketplace_id=apj6jra9ng5v4]#1.height.value": 1.0,
                    "item_package_weight[marketplace_id=apj6jra9ng5v4]#1.value": 119.0,
                    "color[marketplace_id=apj6jra9ng5v4][language_tag=it_it]#1.value": "Nero",
                    "supplier_declared_dg_hz_regulation1": "Nicht zutreffend",
                },
            }
        )

        assert record is not None

        row = processor._build_output_row(
            source_record=record,
            new_color_code="BK",
            new_size_code="00",
            profile=processor_module.COUNTRY_PROFILES[Country.IT],
        )

        self.assertEqual(row["Brand Name"], "Marchio Test")
        self.assertEqual(row["Outer Material Type"], "Poliestere")
        self.assertEqual(row["generic_keywords"], "sottoveste donna")
        self.assertEqual(row["Update Delete"], "Aggiorna")
        self.assertEqual(row["Condition Type"], "Nuovo")
        self.assertEqual(row["Country/Region Of Origin"], "Cina")
        self.assertEqual(row["country_of_origin"], "Cina")
        self.assertEqual(row["pattern_type"], "tinta unita")
        self.assertEqual(row["supplier_declared_material_regulation1"], "Non applicabile")
        self.assertEqual(row["package_level"], "unit")
        self.assertEqual(row["package_weight"], 500.0)
        self.assertEqual(row["package_weight_unit_of_measure"], "GR")
        self.assertEqual(row["package_length_unit_of_measure"], "IN")
        self.assertEqual(row["package_height_unit_of_measure"], "IN")
        self.assertEqual(row["package_width_unit_of_measure"], "IN")
        self.assertEqual(row["Colour"], "Nero")
        self.assertEqual(row["Colour Map"], "Nero")
        self.assertEqual(row["Package Length"], 30.0)
        self.assertEqual(row["Package Width"], 20.0)
        self.assertEqual(row["Package Height"], 1.0)
        self.assertNotIn("Package Weight", row)
        self.assertNotIn("supplier_declared_dg_hz_regulation1", row)

    def test_build_output_row_uses_fr_localized_values_and_aliases(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EA02333BD00",
                "parsed_sku": processor_module.parse_sku("EA02333BD00"),
                "price": 19.9,
                "asin": "B0TESTASIN",
                "all_listings_data": {"item_name": "Fond de robe Bleu Royal 32FR"},
                "category_data": {
                    "brand[marketplace_id=a13v1ib3viyzzh][language_tag=fr_fr]#1.value": "Marque Test",
                    "generic_keyword[marketplace_id=a13v1ib3viyzzh][language_tag=fr_fr]#1.value": (
                        "fond de robe femme"
                    ),
                    "color[marketplace_id=a13v1ib3viyzzh][language_tag=fr_fr]#1.value": "Bleu Royal",
                    "supplier_declared_material_regulation1": "Non applicabile",
                    "supplier_declared_dg_hz_regulation1": "Nicht zutreffend",
                },
            }
        )

        assert record is not None

        row = processor._build_output_row(
            source_record=record,
            new_color_code="BD",
            new_size_code="00",
            profile=processor_module.COUNTRY_PROFILES[Country.FR],
        )

        self.assertEqual(row["Brand Name"], "Marque Test")
        self.assertEqual(row["generic_keywords"], "fond de robe femme")
        self.assertEqual(row["Update Delete"], "Actualisation")
        self.assertEqual(row["Condition Type"], "Neuf")
        self.assertEqual(row["Colour"], "Bleu Royal")
        self.assertEqual(row["Colour Map"], "Bleu")
        self.assertEqual(row["Country/Region Of Origin"], "Chine")
        self.assertEqual(row["country_of_origin"], "Chine")
        self.assertEqual(row["pattern_type"], "Unie")
        self.assertEqual(row["is_adult_product"], "Non")
        self.assertEqual(row["supplier_declared_dg_hz_regulation1"], "Not Applicable")
        self.assertEqual(row["package_weight"], 0.4)
        self.assertEqual(row["package_weight_unit_of_measure"], "GR")
        self.assertEqual(row["package_dimensions_unit_of_measure"], "IN")
        self.assertNotIn("supplier_declared_material_regulation1", row)

    def test_infer_colour_map_localizes_supported_languages(self) -> None:
        processor = AddColorSizeProcessor()

        self.assertEqual(processor._infer_colour_map("Royal Blue", lang="en"), "Blue")
        self.assertEqual(processor._infer_colour_map("Royal Blue", lang="de"), "Blau")
        self.assertEqual(processor._infer_colour_map("Royal Blue", lang="fr"), "Bleu")
        self.assertEqual(processor._infer_colour_map("Royal Blue", lang="it"), "Blu")
        self.assertEqual(processor._infer_colour_map("Royal Blue", lang="es"), "Azul")

    def test_write_workbook_falls_back_to_blank_workbook_when_template_is_missing(self) -> None:
        headers = ["Seller SKU", "Colour", "main_image_url"]
        rows_by_suffix = {
            "-UK1": [{"Seller SKU": "SKU1", "Colour": "Royal Blue", "main_image_url": "https://example.com/1.jpg"}],
            "-UK2": [{"Seller SKU": "SKU2", "Colour": "Sage Green", "main_image_url": "https://example.com/2.jpg"}],
        }
        output_path = Path(self._results_dir.name) / "fallback-output.xlsx"

        AddColorSizeProcessor()._write_workbook(
            output_path=output_path,
            headers=headers,
            rows_by_suffix=rows_by_suffix,
            template_path=processor_module.TEMPLATES_DIR / "missing-template.xlsx",
        )

        workbook = load_workbook(output_path, read_only=True, data_only=True)
        try:
            self.assertEqual(workbook.sheetnames, ["-UK1", "-UK2"])
        finally:
            workbook.close()

        _, rows = _sheet_rows(output_path, "-UK1")
        self.assertEqual(rows[0]["Seller SKU"], "SKU1")
        self.assertEqual(rows[0]["main_image_url"], "https://example.com/1.jpg")

    def test_get_source_colour_name_prefers_localized_display_value_for_item_name_replacement(self) -> None:
        processor = AddColorSizeProcessor()
        record = processor._normalize_record(
            {
                "sku": "EA02333BD00",
                "parsed_sku": processor_module.parse_sku("EA02333BD00"),
                "category_data": {
                    "color": "Rot",
                    "color[marketplace_id=a1pa6795ukmfr9][language_tag=de_de]#1.value": "Burgund",
                },
            }
        )

        assert record is not None

        old_colour = processor._get_source_colour_name(record, lang="de")
        product_name = processor._build_product_name(
            source_data={"item_name": "Unterkleid Burgund 34DE"},
            old_colour=old_colour,
            new_colour="Weiß",
            old_size_map=34,
            new_size_map=34,
            country_code="DE",
        )

        self.assertEqual(old_colour, "Burgund")
        self.assertEqual(product_name, "Unterkleid Weiß 34DE")

    def test_process_request_generates_it_output_with_italian_static_values(self) -> None:
        self._create_it_template_workbook()
        self._require_upload_files(
            "DE_all_listings_20260318062018_所有商品报告_03-18-2026.txt",
            "DE_category_20260318062031_0_APPAREL_HEAD_NECK_COVERING-DRESS.xlsm",
            "DE_category_20260318062034_1_DRESS-UNDERGARMENT_SLIP.xlsm",
        )

        request = ProcessRequest(
            country=Country.IT,
            all_listings_files=["DE_all_listings_20260318062018_所有商品报告_03-18-2026.txt"],
            category_files=[
                "DE_category_20260318062031_0_APPAREL_HEAD_NECK_COVERING-DRESS.xlsm",
                "DE_category_20260318062034_1_DRESS-UNDERGARMENT_SLIP.xlsm",
            ],
            selected_prefixes=["EA02333"],
            target_colors=["BK"],
            start_size="00",
            end_size="00",
            size_step=1,
            mode="add-color",
        )

        result = AddColorSizeProcessor().process(request, output_filename="it-request-output")

        self.assertEqual(result.output_file, "it-request-output.xlsx")
        self.assertEqual(result.processed_count, 1)
        self.assertEqual(result.skipped_count, 0)
        self.assertEqual(result.errors, [])

        output_path = Path(self._results_dir.name) / result.output_file
        _, rows = _template_machine_rows(output_path)
        children = self._child_rows(rows)

        self.assertEqual(len(rows), 2)
        self.assertEqual(len(children), 1)
        self.assertEqual(children[0]["item_sku"], "EA02333BK00")
        self.assertEqual(children[0]["update_delete"], "Aggiorna")
        self.assertEqual(children[0]["condition_type"], "Nuovo")
        self.assertEqual(children[0]["color_name"], "Nero")
        self.assertEqual(children[0]["color_map"], "Nero")
        self.assertEqual(children[0]["country_of_origin"], "Cina")
        self.assertEqual(children[0]["pattern_type"], "tinta unita")
        self.assertEqual(children[0]["supplier_declared_material_regulation1"], "Non applicabile")
        self.assertEqual(children[0]["package_level"], "unit")
        self.assertEqual(children[0]["package_weight"], 500.0)
        self.assertEqual(children[0]["package_weight_unit_of_measure"], "GR")
        self.assertEqual(children[0]["package_length_unit_of_measure"], "IN")
        self.assertEqual(children[0]["package_height_unit_of_measure"], "IN")
        self.assertEqual(children[0]["package_width_unit_of_measure"], "IN")
        self.assertIsNone(children[0]["supplier_declared_dg_hz_regulation1"])

    def test_write_workbook_preserves_template_row_styles_for_new_rows(self) -> None:
        template_path = self._create_template_workbook()
        workbook = load_workbook(template_path)
        try:
            worksheet = workbook["Template"]
            orange_fill = PatternFill(fill_type="solid", fgColor="F4B183")
            bold_font = Font(name="Calibri", bold=True, color="9C0006")
            border = Border(left=Side(style="thin", color="C65911"), right=Side(style="thin", color="C65911"))
            alignment = Alignment(horizontal="center")

            styled_columns = (1, 2, 3, 4)
            for col_idx in styled_columns:
                cell = worksheet.cell(row=4, column=col_idx)
                cell.fill = orange_fill
                cell.font = bold_font
                cell.border = border
                cell.alignment = alignment
                cell.number_format = "@"

            workbook.save(template_path)
        finally:
            workbook.close()

        output_path = Path(self._results_dir.name) / "styled-output.xlsx"
        AddColorSizeProcessor()._write_workbook(
            output_path=output_path,
            headers=["Seller SKU", "Parent SKU", "Variation Theme", "Colour", "Colour Map"],
            rows_by_suffix={
                "-UK1": [
                    {
                        "Seller SKU": "SKU1",
                        "Parent SKU": "PARENT1",
                        "Variation Theme": "SizeName-ColorName",
                        "Colour": "Royal Blue",
                        "Colour Map": "Blue",
                    }
                ]
            },
            template_path=template_path,
        )

        styled_workbook = load_workbook(output_path)
        try:
            worksheet = styled_workbook["Template"]
            self.assertEqual(worksheet["A4"].fill.fgColor.rgb, "00F4B183")
            self.assertTrue(worksheet["A4"].font.bold)
            self.assertEqual(worksheet["A4"].alignment.horizontal, "center")
            self.assertEqual(worksheet["A4"].number_format, "@")
            self.assertEqual(worksheet["E4"].fill.fgColor.rgb, "00FCE4D6")
        finally:
            styled_workbook.close()


if __name__ == "__main__":
    unittest.main()
