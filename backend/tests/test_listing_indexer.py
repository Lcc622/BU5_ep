from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
import sys

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.indexers.listing_indexer import ListingIndexer


class ListingIndexerTests(unittest.TestCase):
    def test_build_index_merges_reports_and_keeps_unparsed_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            temp_path = Path(tmp_dir)
            listings_file = temp_path / "all_listings.txt"
            category_file = temp_path / "category.xlsx"

            listings_file.write_text(
                "\ufeffseller-sku\titem-name\tprice\tasin1\tparent-asin\tlisting-id\tstatus\n"
                "AB12345BK01-UK\tListing Name\t19.99\tASIN1\tPARENT1\tLIST1\tActive\n"
                "bad sku\tBroken Item\t9.99\tASIN2\tPARENT2\tLIST2\tInactive\n",
                encoding="utf-8",
            )

            workbook = Workbook()
            template_sheet = workbook.active
            template_sheet.title = "Template"
            template_sheet.append(["ignored"])
            template_sheet.append(["SKU", "Item Name", "Colour", "bullet_points", "price"])
            template_sheet.append(["AB12345BK01-UK", "Category Name", "Black", "soft fabric", "25.00"])
            template_sheet.append(["bad sku", "Broken Category", "Gray", "fallback bullet", "10.00"])
            workbook.save(category_file)

            index = ListingIndexer().build_index([listings_file], [category_file])

        sku_record = index["by_sku"]["AB12345BK01-UK"]
        self.assertEqual(sku_record["prefix"], "AB12345")
        self.assertEqual(sku_record["merged_data"]["item name"], "Category Name")
        self.assertEqual(sku_record["merged_data"]["price"], "25.00")
        self.assertEqual(sku_record["all_listings_data"]["listing-id"], "LIST1")
        self.assertEqual(sku_record["category_data"]["colour"], "Black")
        self.assertIn("AB12345BK01-UK", index["by_prefix"]["AB12345"])

        broken_record = index["by_sku"]["BAD SKU"]
        self.assertIsNone(broken_record["parsed_sku"])
        self.assertEqual(index["stats"]["unparsed_count"], 2)

    def test_build_index_falls_back_to_first_sheet_with_sku_column(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            temp_path = Path(tmp_dir)
            listings_file = temp_path / "all_listings.txt"
            category_file = temp_path / "category.xlsx"

            listings_file.write_text(
                "seller-sku\titem-name\tprice\n"
                "AB12345BK01-UK\tListing Name\t19.99\n",
                encoding="utf-8",
            )

            workbook = Workbook()
            workbook.active.title = "Intro"
            data_sheet = workbook.create_sheet("Data Export")
            data_sheet.append(["Notes"])
            data_sheet.append(["SKU", "Item Name", "Colour Map"])
            data_sheet.append(["AB12345BK01-UK", "Sheet Name", "black"])
            workbook.save(category_file)

            index = ListingIndexer().build_index([listings_file], [category_file])

        sku_record = index["by_sku"]["AB12345BK01-UK"]
        self.assertEqual(sku_record["category_data"]["colour map"], "black")

    def test_build_index_prefers_machine_header_row_on_lowercase_template_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmp_dir:
            temp_path = Path(tmp_dir)
            listings_file = temp_path / "all_listings.txt"
            category_file = temp_path / "category.xlsx"

            listings_file.write_text(
                "seller-sku\titem-name\tprice\n"
                "AB12345BK01-UK\tListing Name\t19.99\n",
                encoding="utf-8",
            )

            workbook = Workbook()
            export_sheet = workbook.active
            export_sheet.title = "Data Export"
            export_sheet.append(["SKU", "Item Name"])
            export_sheet.append(["AB12345BK01-UK", "Wrong Sheet"])

            template_sheet = workbook.create_sheet("template")
            template_sheet.append(["meta"])
            template_sheet.append(["meta"])
            template_sheet.append(["meta"])
            template_sheet.append(["SKU", "Markenname", "Artikelname", "Pflegehinweise"])
            template_sheet.append(
                [
                    "contribution_sku#1.value",
                    "brand[marketplace_id=A1PA6795UKMFR9][language_tag=de_DE]#1.value",
                    "item_name[marketplace_id=A1PA6795UKMFR9][language_tag=de_DE]#1.value",
                    "care_instructions[marketplace_id=A1PA6795UKMFR9][language_tag=de_DE]#1.value",
                ]
            )
            template_sheet.append(["AB12345BK01-UK", "Ever-Pretty", "Abendkleid", "Hand Wash Only"])
            workbook.save(category_file)

            index = ListingIndexer().build_index([listings_file], [category_file])

        sku_record = index["by_sku"]["AB12345BK01-UK"]
        self.assertEqual(sku_record["category_data"]["sku"], "AB12345BK01-UK")
        self.assertEqual(sku_record["category_data"]["brand"], "Ever-Pretty")
        self.assertEqual(sku_record["category_data"]["item_name"], "Abendkleid")
        self.assertEqual(sku_record["category_data"]["care_instructions"], "Hand Wash Only")
        self.assertEqual(
            sku_record["category_data"][
                "brand[marketplace_id=a1pa6795ukmfr9][language_tag=de_de]#1.value"
            ],
            "Ever-Pretty",
        )


if __name__ == "__main__":
    unittest.main()
